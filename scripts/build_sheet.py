# -*- coding: utf-8 -*-
"""VDO Script DP Studio: รวม shots.json + analysis.json เป็นตาราง Script

ใช้: python3 build_sheet.py <work_dir>
ต้องมี <work_dir>/shots.json (จาก extract.py) และ <work_dir>/analysis.json (Claude เขียน)
ได้ไฟล์: script.md, script.csv, script.xlsx, script.html (หน้าเว็บ) และ script.docx (ถ้ามี python-docx)
ถ้า analysis.json มีบทดัดแปลง (adapt_vo) จะเพิ่มคอลัมน์บทดัดแปลงพร้อมนับคำเทียบต้นฉบับ
"""
import base64, csv, html, io, json, os, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wordcount import TOLERANCE, badge_class, count_words  # noqa: E402

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

COLS = [('n', '#'), ('time', 'Sequence / Time'), ('shot', 'Shot Size'), ('angle', 'Angle'),
        ('movement', 'Camera Movement'), ('description', 'Description (ภาพ)'),
        ('vo', 'Voiceover / Dialogue'), ('text_on_screen', 'Text on Screen'),
        ('audio', 'Music / SFX'), ('notes', 'Notes')]
EMPTY = ('', '-')


def cell(v):
    return str(v or '').replace('|', '/').replace('\n', '<br>')


def esc(v):
    return html.escape(str(v or ''))


def thumb_b64(path, size=(216, 384)):
    from PIL import Image
    im = Image.open(path).convert('RGB'); im.thumbnail(size)
    b = io.BytesIO(); im.save(b, 'JPEG', quality=72)
    return base64.b64encode(b.getvalue()).decode()


def load(work):
    S = json.load(io.open(os.path.join(work, 'shots.json'), encoding='utf-8'))
    A = json.load(io.open(os.path.join(work, 'analysis.json'), encoding='utf-8'))
    by_n = {int(x['n']): x for x in A.get('shots', [])}
    adapt = A.get('adapt') or {}
    has_adapt = any((x.get('adapt_vo') or '').strip() for x in A.get('shots', []))
    brand = adapt.get('brand') or 'DP Studio'
    rows = []
    for s in S['shots']:
        a = by_n.get(s['n'], {})
        r = {'n': s['n'], 'time': s['time'], 'frame': s['frame'], 'dur': round(s['end'] - s['start'], 1)}
        for k in ('shot', 'angle', 'movement', 'description', 'vo', 'vo_th', 'text_on_screen', 'audio', 'notes',
                  'adapt_visual', 'adapt_vo', 'adapt_text'):
            r[k] = a.get(k, '')
        if not r['vo']:
            r['vo'] = s.get('vo_auto', '')
        r['wc_src'] = count_words(r['vo'])
        r['wc_adapt'] = count_words(r['adapt_vo']) if has_adapt else 0
        r['wc_diff'] = r['wc_adapt'] - r['wc_src']
        rows.append(r)
    return S, A, rows, has_adapt, brand, adapt


def main():
    work = sys.argv[1]
    S, A, rows, has_adapt, brand, adapt = load(work)
    has_th = any(r.get('vo_th') for r in rows)
    cols = list(COLS)
    if has_th:
        cols.insert(7, ('vo_th', 'แปลไทย'))
    cols.insert(cols.index(('vo', 'Voiceover / Dialogue')) + 1, ('wc_src', 'คำ'))
    if has_adapt:
        cols += [('adapt_visual', 'ภาพ %s' % brand), ('adapt_vo', 'บทพูด %s' % brand), ('wc_adapt', 'คำ %s' % brand),
                 ('wc_diff', 'ส่วนต่าง'), ('adapt_text', 'บนจอ %s' % brand)]
    tot_src = sum(r['wc_src'] for r in rows)
    tot_adapt = sum(r['wc_adapt'] for r in rows)

    title = A.get('title') or S['meta'].get('title') or 'Video'
    summ = A.get('summary', {})
    ad = A.get('ad_copy') or {}
    if not ad.get('en') and S['meta'].get('ad_text'):
        ad['en'] = S['meta']['ad_text']
    info = [('ชื่อวิดีโอ', title), ('ที่มา', S['source']), ('ความยาว', '%.1f วินาที' % S['duration']),
            ('สัดส่วนภาพ', '%s (%sx%s, %s fps)' % (S['aspect'], S['width'], S['height'], S['fps'])),
            ('จำนวนช็อต', '%d ช็อต (เฉลี่ย %.1f วินาที/ช็อต)' % (S['shot_count'], S['avg_shot'])),
            ('ถอดเสียงจาก', S['transcript_source']), ('จำนวนคำต้นฉบับ', '%d คำ' % tot_src)]
    if has_adapt:
        info.append(('จำนวนคำบท %s' % brand, '%d คำ (ต่างกัน %+d)' % (tot_adapt, tot_adapt - tot_src)))
    labels = {'hook': 'Hook (3 วินาทีแรก)', 'structure': 'โครงเรื่อง', 'style': 'สไตล์ภาพและการตัดต่อ',
              'cta': 'CTA', 'takeaways': 'เอาไปใช้ต่อได้', 'notes': 'หมายเหตุ'}
    info += [(labels.get(k, k), v) for k, v in summ.items() if v]
    for k, h in (('en', 'แคปชันแอดต้นฉบับ'), ('th', 'แคปชันแปลไทย'), ('cta', 'ปุ่ม CTA ต้นฉบับ'),
                 ('adapt', 'แคปชัน %s' % brand), ('cta_adapt', 'ปุ่ม CTA %s' % brand)):
        if ad.get(k):
            info.append((h, ad[k]))

    # Markdown
    md = ['# VDO Script: %s' % title, '', '| หัวข้อ | รายละเอียด |', '|---|---|']
    md += ['| %s | %s |' % (k, cell(v)) for k, v in info]
    md += ['', '## Shot List', '', '| ' + ' | '.join(h for _, h in cols) + ' |', '|' + '---|' * len(cols)]
    for r in rows:
        md.append('| ' + ' | '.join(cell(r.get(k)) for k, _ in cols) + ' |')
    if S.get('full_transcript'):
        md += ['', '## Transcript ทั้งหมด', '', S['full_transcript']]
    md += ['', '_สร้างด้วย VDO Script DP Studio · ใช้เพื่อศึกษาและอ้างอิงเท่านั้น ลิขสิทธิ์วิดีโอเป็นของเจ้าของผลงาน_']
    io.open(os.path.join(work, 'script.md'), 'w', encoding='utf-8').write('\n'.join(md))

    # CSV (เปิดใน Excel / Google Sheets ภาษาไทยไม่เพี้ยน)
    with io.open(os.path.join(work, 'script.csv'), 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow([h for _, h in cols])
        for r in rows:
            w.writerow([r.get(k, '') for k, _ in cols])

    # Excel พร้อมภาพตัวอย่างทุกช็อต
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImg
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from PIL import Image
    wb = Workbook()
    ws = wb.active; ws.title = 'Shot List'
    head = PatternFill('solid', fgColor='111820')
    tint = PatternFill('solid', fgColor='EEF5FB')
    xcols = cols[:2] + [('frame', 'Frame')] + cols[2:]
    widths = {'n': 5, 'time': 14, 'frame': 26, 'shot': 11, 'angle': 13, 'movement': 14, 'description': 40,
              'vo': 40, 'vo_th': 36, 'wc_src': 6, 'text_on_screen': 28, 'audio': 20, 'notes': 24,
              'adapt_visual': 30, 'adapt_vo': 40, 'wc_adapt': 8, 'wc_diff': 8, 'adapt_text': 26}
    for c, (k, h) in enumerate(xcols, 1):
        x = ws.cell(1, c, h); x.fill = head; x.font = Font(bold=True, color='F8B93E')
        x.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths.get(k, 16)
    thumbs = os.path.join(work, 'frames', '_xlsx'); os.makedirs(thumbs, exist_ok=True)
    for i, r in enumerate(rows, 2):
        ws.row_dimensions[i].height = 100
        for c, (k, _) in enumerate(xcols, 1):
            if k == 'frame':
                if os.path.exists(r['frame']):
                    im = Image.open(r['frame']); im.thumbnail((180, 128))
                    tp = os.path.join(thumbs, os.path.basename(r['frame']).replace('.jpg', '.png'))
                    im.save(tp)
                    ws.add_image(XImg(tp), '%s%d' % (get_column_letter(c), i))
                continue
            x = ws.cell(i, c, r.get(k, ''))
            x.alignment = Alignment(vertical='top', wrap_text=True)
            if k.startswith('adapt') or k in ('wc_adapt', 'wc_diff'):
                x.fill = tint
            if k == 'wc_diff' and r['wc_src'] and abs(r['wc_diff']) > TOLERANCE:
                x.font = Font(bold=True, color='C23D68')
    ws.freeze_panes = 'D2'
    ws2 = wb.create_sheet('Summary')
    ws2.column_dimensions['A'].width = 26; ws2.column_dimensions['B'].width = 100
    for i, (k, v) in enumerate(info, 1):
        ws2.cell(i, 1, k).font = Font(bold=True)
        ws2.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical='top')
    if S.get('full_transcript'):
        n = len(info) + 2
        ws2.cell(n, 1, 'Transcript ทั้งหมด').font = Font(bold=True)
        ws2.cell(n, 2, S['full_transcript']).alignment = Alignment(wrap_text=True, vertical='top')
    xp = os.path.join(work, 'script.xlsx')
    wb.save(xp)

    hp = build_html(work, S, A, rows, has_adapt, brand, adapt, ad, title, summ, tot_src, tot_adapt)
    dp = build_docx(work, S, rows, has_adapt, brand, ad, title, tot_src, tot_adapt)

    off = [r['n'] for r in rows if has_adapt and r['wc_src'] and abs(r['wc_diff']) > TOLERANCE]
    print(json.dumps({'markdown': os.path.join(work, 'script.md'), 'xlsx': xp, 'csv': os.path.join(work, 'script.csv'),
                      'html': hp, 'docx': dp, 'rows': len(rows), 'words_source': tot_src,
                      'words_adapt': tot_adapt if has_adapt else None, 'shots_word_count_off': off},
                     ensure_ascii=False, indent=1))


def build_html(work, S, A, rows, has_adapt, brand, adapt, ad, title, summ, tot_src, tot_adapt):
    tpl = io.open(os.path.join(SKILL_DIR, 'assets', 'page.html'), encoding='utf-8').read()

    def badge(c, d=None):
        if d is None:
            return '<span class="wc">%d คำ</span>' % c
        return '<span class="wc %s">%d คำ · %+d</span>' % (badge_class(d), c, d)

    trs = []
    for r in rows:
        img = ('<img src="data:image/jpeg;base64,%s" alt="ช็อต %d ต้นฉบับ" width="108" height="192" loading="lazy">'
               % (thumb_b64(r['frame']), r['n'])) if os.path.exists(r['frame']) else ''
        chips = ''.join('<span class="chip">%s</span>' % esc(r[k]) for k in ('shot', 'angle', 'movement') if r.get(k))
        vo = r['vo'] if r['vo'] not in EMPTY else '-'
        orig = '<p class="desc">%s</p><p class="vo en">%s %s</p>' % (esc(r['description']), esc(vo), badge(r['wc_src']) if r['wc_src'] else '')
        if r.get('vo_th') and r['vo_th'] not in EMPTY:
            orig += '<p class="vo th"><span class="lb">แปลไทย</span>%s</p>' % esc(r['vo_th'])
        orig += '<p class="ost"><span>บนจอ</span>%s</p>' % esc(r['text_on_screen'] or '-')
        if r.get('notes'):
            orig += '<p class="note">%s</p>' % esc(r['notes'])
        dpc = ''
        if has_adapt:
            avo = r['adapt_vo'] if r['adapt_vo'] not in EMPTY else '(ไม่มีเสียงพูด)'
            dpc = '<p class="desc">%s</p><p class="vo dpvo">%s %s</p><p class="ost"><span>บนจอ</span>%s</p>' % (
                esc(r['adapt_visual']), esc(avo), badge(r['wc_adapt'], r['wc_diff']) if r['wc_src'] else '', esc(r['adapt_text'] or '-'))
        trs.append('<tr id="s%d"><td class="c-n"><span class="num">%d</span></td><td class="c-img">%s</td>'
                   '<td class="c-t"><span class="tc">%s</span><span class="dur">%.1f วิ</span></td>'
                   '<td class="c-cam">%s</td><td class="col-orig">%s</td><td class="col-dp">%s</td></tr>'
                   % (r['n'], r['n'], img, esc(r['time'].split('-')[0]), r['dur'], chips, orig, dpc))

    def mmss(sec):
        return '%d:%02d' % (int(sec) // 60, round(sec) % 60)

    meta = ''.join('<span>%s <b>%s</b></span>' % (k, esc(v)) for k, v in (
        ('ความยาว', mmss(S['duration'])), ('สัดส่วน', S['aspect']), ('ช็อต', S['shot_count']),
        ('เฉลี่ย', '%.2f วิ/ช็อต' % S['avg_shot'])))
    if S['source'].startswith('http'):
        meta += '<span><a href="%s" target="_blank" rel="noopener">ดูวิดีโอต้นฉบับ</a></span>' % esc(S['source'])
    sm_labels = (('hook', 'Hook ต้นฉบับ'), ('structure', 'โครงเรื่อง'), ('style', 'สไตล์ภาพ'))
    summary = ''.join('<div><h3>%s</h3><p>%s</p></div>' % (h, esc(summ.get(k))) for k, h in sm_labels if summ.get(k))

    brief = ''
    if ad.get('en'):
        cells = ['<div><h3>ต้นฉบับ</h3><p>%s</p>%s</div>' % (esc(ad['en']), ('<p class="cta">ปุ่ม <b>%s</b></p>' % esc(ad['cta'])) if ad.get('cta') else '')]
        if ad.get('th'):
            cells.append('<div><h3>แปลไทย</h3><p>%s</p></div>' % esc(ad['th']))
        if ad.get('adapt'):
            cells.append('<div class="dp"><h3>แคปชัน %s</h3><p>%s</p>%s</div>' % (
                esc(brand), esc(ad['adapt']), ('<p class="cta">ปุ่ม <b>%s</b></p>' % esc(ad['cta_adapt'])) if ad.get('cta_adapt') else ''))
        brief = ('<section style="display:grid;gap:12px" aria-label="บรีฟแคปชัน"><div class="sec-h"><h2>บรีฟจากแคปชันแอด</h2>'
                 '<p>ข้อความ Primary text ที่แสดงเหนือวิดีโอ</p></div><div class="brief%s">%s</div></section>'
                 % (' two' if len(cells) == 2 else '', ''.join(cells)))

    count = '<span><b>%d</b>คำ ต้นฉบับ</span>' % tot_src
    if has_adapt:
        count += '<span class="dpc"><b>%d</b>คำ บท %s</span><span>ป้ายเขียว = ห่างจากต้นฉบับไม่เกิน %d คำ</span>' % (tot_adapt, esc(brand), TOLERANCE)

    notes = ''
    secs = [(s.get('heading', ''), s.get('items', [])) for s in adapt.get('sections', [])]
    if secs:
        notes = '<div class="notes">' + ''.join('<section><h2>%s</h2><ul>%s</ul></section>' % (
            esc(h), ''.join('<li>%s</li>' % esc(i) for i in items)) for h, items in secs) + '</div>'

    # ตารางสะอาดสำหรับวางใน Google Docs
    TD = 'style="border:1px solid #999;padding:6px;vertical-align:top;font-family:Prompt,sans-serif;font-size:10pt"'
    TH = 'style="border:1px solid #999;padding:6px;background:#111820;color:#ffffff;font-family:Prompt,sans-serif;font-size:10pt;text-align:left"'
    heads = ['#', 'เวลา', 'ภาพ', 'บทพูดต้นฉบับ', 'แปลไทย', 'ตัวหนังสือบนจอ']
    if has_adapt:
        heads = ['#', 'เวลา', 'ภาพ %s' % brand, 'ต้นฉบับ', 'แปลไทย', 'บทพูด %s' % brand, 'บนจอ %s' % brand]
    crow = []
    for r in rows:
        src = (r['vo'] + (' (%d คำ)' % r['wc_src'] if r['wc_src'] else '')) if r['vo'] not in EMPTY else '-'
        if has_adapt:
            avo = (r['adapt_vo'] + ' (%d คำ)' % r['wc_adapt']) if r['adapt_vo'] not in EMPTY else '-'
            cs = [r['n'], r['time'], r['adapt_visual'], src, r.get('vo_th', ''), avo, r['adapt_text']]
        else:
            cs = [r['n'], r['time'], r['description'], src, r.get('vo_th', ''), r['text_on_screen']]
        crow.append('<tr>' + ''.join('<td %s>%s</td>' % (TD, esc(c)) for c in cs) + '</tr>')
    h1_text = adapt.get('title') if has_adapt and adapt.get('title') else title
    intro = '<p style="font-family:Prompt,sans-serif;font-size:14pt;font-weight:bold">%s</p>' % esc(h1_text)
    lines = []
    if has_adapt and ad.get('adapt'):
        lines.append('แคปชัน: ' + ad['adapt'])
    lines.append('จำนวนคำ: ต้นฉบับ %d คำ%s' % (tot_src, (' · บท %s %d คำ' % (brand, tot_adapt)) if has_adapt else ''))
    intro += '<p style="font-family:Prompt,sans-serif;font-size:10pt">%s</p>' % '<br>'.join(esc(x) for x in lines)
    copytbl = intro + '<table style="border-collapse:collapse;width:100%"><tr>' + ''.join('<th %s>%s</th>' % (TH, h) for h in heads) + '</tr>' + ''.join(crow) + '</table>'

    key = 'adapt_vo' if has_adapt else 'vo'
    plain = '\n'.join('%d. [%s] %s' % (r['n'], r['time'].split('-')[0], r[key]) for r in rows if r[key] not in EMPTY)

    if has_adapt:
        h1 = esc(h1_text)
        lede = esc(adapt.get('lede') or 'สคริปต์ %s ดัดแปลงจาก %s วางภาพต้นฉบับกับบทของเราไว้ช็อตต่อช็อต' % (brand, title))
        eyebrow = '%s · Ad Script' % esc(brand)
    else:
        h1 = esc(title)
        lede = esc(summ.get('hook', ''))
        eyebrow = 'VDO Script'
    rep = {
        '{{TITLE}}': esc(h1_text)[:60], '{{BODYCLASS}}': '' if has_adapt else 'no-adapt',
        '{{EYEBROW}}': eyebrow, '{{H1}}': h1, '{{LEDE}}': lede, '{{META}}': meta, '{{SUMMARY}}': summary,
        '{{BRIEF}}': brief, '{{COUNT}}': count, '{{ADAPT_SHORT}}': 'บท %s' % esc(brand),
        '{{COPY_LABEL}}': ('คัดลอกบทพูด %s' % esc(brand)) if has_adapt else 'คัดลอกบทพูด',
        '{{SRC_HEAD}}': 'ต้นฉบับ + แปลไทย' if any(r.get('vo_th') for r in rows) else 'ต้นฉบับ',
        '{{ADAPT_HEAD}}': 'เวอร์ชัน %s' % esc(brand), '{{ROWS}}': '\n'.join(trs), '{{NOTES}}': notes,
        '{{FOOT}}': 'ภาพในตารางเป็นเฟรมจากวิดีโอต้นฉบับ ใช้เพื่ออ้างอิงภายในทีมเท่านั้น ลิขสิทธิ์เป็นของเจ้าของผลงาน · สร้างด้วย /vdo-script',
        '{{COPYTBL}}': copytbl, '{{DPTEXT}}': json.dumps(plain, ensure_ascii=False).replace('</', '<\\/'),
    }
    for k, v in rep.items():
        tpl = tpl.replace(k, str(v))
    hp = os.path.join(work, 'script.html')
    io.open(hp, 'w', encoding='utf-8').write(tpl)
    return hp


def build_docx(work, S, rows, has_adapt, brand, ad, title, tot_src, tot_adapt):
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        return None
    from PIL import Image
    doc = Document()
    sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, m, Cm(1.5))
    st = doc.styles['Normal']; st.font.name = 'Prompt'; st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn('w:cs'), 'Prompt')

    def run(p, t, b=False, size=None, color=None):
        r = p.add_run(str(t)); r.bold = b; r.font.name = 'Prompt'
        r._element.rPr.rFonts.set(qn('w:cs'), 'Prompt')
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return r

    run(doc.add_paragraph(), 'VDO Script: %s' % title, True, 16)
    run(doc.add_paragraph(), 'ที่มา: %s · %.0f วินาที · %s · %d ช็อต' % (S['source'], S['duration'], S['aspect'], S['shot_count']), size=9, color='5A6270')
    for k, h in (('en', 'แคปชันต้นฉบับ'), ('th', 'แปลไทย'), ('adapt', 'แคปชัน %s' % brand)):
        if ad.get(k):
            p = doc.add_paragraph(); run(p, h + ': ', True); run(p, ad[k])
    if has_adapt:
        heads = ['#', 'ภาพต้นฉบับ', 'เวลา', 'ต้นฉบับ', 'แปลไทย', 'บทพูด %s' % brand, 'ภาพ / บนจอ %s' % brand]
        widths = [0.8, 2.4, 1.8, 5.5, 5.0, 5.5, 5.7]
    else:
        heads = ['#', 'ภาพต้นฉบับ', 'เวลา', 'ภาพ', 'บทพูด', 'แปลไทย', 'บนจอ']
        widths = [0.8, 2.4, 1.8, 5.5, 5.5, 5.0, 5.7]
    t = doc.add_table(rows=1, cols=len(heads)); t.style = 'Table Grid'
    for i, h in enumerate(heads):
        c = t.rows[0].cells[i]; c.width = Cm(widths[i]); run(c.paragraphs[0], h, True, color='FFFFFF')
        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), '111820')
        c._tc.get_or_add_tcPr().append(sh)
    for r in rows:
        cells = t.add_row().cells
        for i, c in enumerate(cells):
            c.width = Cm(widths[i])
        run(cells[0].paragraphs[0], r['n'])
        if os.path.exists(r['frame']):
            im = Image.open(r['frame']).convert('RGB'); im.thumbnail((200, 356))
            bio = io.BytesIO(); im.save(bio, 'JPEG', quality=80); bio.seek(0)
            cells[1].paragraphs[0].add_run().add_picture(bio, width=Cm(2.0))
        run(cells[2].paragraphs[0], r['time'].split('-')[0]); run(cells[2].add_paragraph(), '%.1f วิ' % r['dur'], size=8, color='5A6270')
        vo = r['vo'] if r['vo'] not in EMPTY else '-'
        if has_adapt:
            run(cells[3].paragraphs[0], vo)
            if r['wc_src']:
                run(cells[3].add_paragraph(), '%d คำ' % r['wc_src'], size=8, color='5A6270')
            run(cells[4].paragraphs[0], r.get('vo_th') or '')
            run(cells[5].paragraphs[0], r['adapt_vo'] if r['adapt_vo'] not in EMPTY else '-', True)
            if r['wc_src']:
                run(cells[5].add_paragraph(), '%d คำ (%+d)' % (r['wc_adapt'], r['wc_diff']), size=8,
                    color='2E7D4F' if abs(r['wc_diff']) <= TOLERANCE else 'C23D68')
            run(cells[6].paragraphs[0], r['adapt_visual']); run(cells[6].add_paragraph(), 'บนจอ: %s' % (r['adapt_text'] or '-'), size=8.5, color='C23D68')
        else:
            run(cells[3].paragraphs[0], r['description'])
            run(cells[4].paragraphs[0], vo)
            run(cells[5].paragraphs[0], r.get('vo_th') or '')
            run(cells[6].paragraphs[0], r['text_on_screen'] or '-')
    p = doc.add_paragraph()
    run(p, 'จำนวนคำรวม: ต้นฉบับ %d คำ%s' % (tot_src, (' · บท %s %d คำ' % (brand, tot_adapt)) if has_adapt else ''), True)
    out = os.path.join(work, 'script.docx')
    doc.save(out)
    return out


if __name__ == '__main__':
    main()
